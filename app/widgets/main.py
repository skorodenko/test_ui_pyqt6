import docx
import numpy as np
from pathlib import Path
from attrs import define, field
from PyQt6.QtWidgets import QMainWindow, QFileDialog
from app.ui.ui_main import Ui_MainWindow
from app.models.table_model import TableModel
from .dialog_expand_table import DialogExpandTable


@define
class MainWindow(QMainWindow, Ui_MainWindow):
    table: np.ndarray = field()
    table_model: TableModel = field()

    @table.default
    def _table_init(self):
        # Change datatype here to contain more then 64 characters
        table = np.full((1, 1), "", dtype=np.dtype("U64"))
        return table

    @table_model.default
    def _table_model_init(self):
        return TableModel(self.table)

    def __attrs_pre_init__(self):
        super().__init__()

    def __attrs_post_init__(self):
        self.setupUi(self)

        self.tableView.setModel(self.table_model)

        self.actionSave_as.triggered.connect(self.saveAs)
        self.actionAdd_One_2.triggered.connect(self.addRow)
        self.actionAdd_One.triggered.connect(self.addColumn)

        self.show()

    def add_row(self, n=1):
        if n == 0:
            self.table_model.update_table(self.table)
            return None
        new_col = np.full((1, self.table.shape[1]), "")
        self.table = np.vstack([self.table, new_col])
        self.add_column(n=n - 1)

    def add_column(self, n=1):
        if n == 0:
            self.table_model.update_table(self.table)
            return None
        new_row = np.full((self.table.shape[0], 1), "")
        self.table = np.hstack([self.table, new_row])
        self.add_row(n=n - 1)

    def addColumn(self):
        dlg = DialogExpandTable("Add columns", parent=self)
        if dlg.exec():
            n = dlg.spinBox.value()
            self.add_column(n)

    def addRow(self):
        dlg = DialogExpandTable("Add rows", parent=self)
        if dlg.exec():
            n = dlg.spinBox.value()
            self.add_row(n)

    def save_to_docx(self, filepath):
        table = self.table
        doc = docx.Document()
        doc_table = doc.add_table(table.shape[0], table.shape[1])

        for i in range(table.shape[0]):
            for j in range(table.shape[1]):
                doc_table.cell(i, j).text = str(table[i, j])

        doc.save(filepath)

    def saveAs(self):
        filename, name_filter = QFileDialog.getSaveFileName(
            self, "Save Table As", filter="Docx (*.docx)"
        )

        filepath = Path(filename)

        match name_filter:
            case "Docx (*.docx)":
                filepath = filepath.with_suffix(".docx")
                self.save_to_docx(filepath)
